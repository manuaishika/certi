"""Registration intake, validation and bulk CSV/XLSX import."""
import csv
import io
import re

from fastapi import HTTPException
from sqlmodel import Session, select

from ..models import Event, Registration
from ..security import new_token

HEADER_ALIASES = {
    "name": "name_en", "name_en": "name_en", "full name": "name_en", "name (english)": "name_en",
    "name_hi": "name_hi", "hindi name": "name_hi", "name (hindi)": "name_hi", "नाम": "name_hi",
    "institution": "institution", "school": "institution", "college": "institution",
    "grade": "grade", "class": "grade", "mobile": "mobile", "phone": "mobile", "mobile number": "mobile",
    "email": "email", "parent": "parent_name", "parent_name": "parent_name", "parent name": "parent_name",
}


def normalize_mobile(raw: str) -> str:
    digits = re.sub(r"\D", "", raw or "")
    if len(digits) > 10 and digits.startswith("91"):
        digits = digits[2:]
    if len(digits) == 11 and digits.startswith("0"):
        digits = digits[1:]
    return digits


def clean_name(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").strip())


def validate(fields: dict) -> list[str]:
    errs = []
    if len(clean_name(fields.get("name_en", ""))) < 2:
        errs.append("Name (English) is required")
    if not re.fullmatch(r"[6-9]\d{9}", normalize_mobile(fields.get("mobile", ""))):
        errs.append("Enter a valid 10-digit mobile number")
    if fields.get("email") and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", fields["email"].strip()):
        errs.append("Email looks invalid")
    hi = fields.get("name_hi", "")
    if hi and not any("ऀ" <= c <= "ॿ" for c in hi):
        errs.append("Devanagari name must be written in Devanagari script (or left blank)")
    return errs


def create_registration(session: Session, event: Event, fields: dict, *, approved: bool | None = None) -> Registration:
    errs = validate(fields)
    if errs:
        raise HTTPException(422, "; ".join(errs))
    mobile = normalize_mobile(fields["mobile"])
    name = clean_name(fields["name_en"])
    dup = session.exec(select(Registration).where(Registration.event_id == event.id,
                                                  Registration.mobile == mobile,
                                                  Registration.name_en == name)).first()
    if dup:
        return dup
    mode = fields.get("attend_mode") or ("online" if event.mode == "online" else "offline")
    if event.mode == "offline":
        mode = "offline"
    reg = Registration(
        event_id=event.id, name_en=name, name_hi=clean_name(fields.get("name_hi", "")),
        parent_name=clean_name(fields.get("parent_name", "")), institution=clean_name(fields.get("institution", "")),
        grade=clean_name(fields.get("grade", "")), mobile=mobile, email=(fields.get("email") or "").strip(),
        attend_mode=mode, token=new_token("P", 10),
        approved=(not event.approval_required) if approved is None else approved)
    session.add(reg)
    session.commit()
    session.refresh(reg)
    return reg


def parse_upload(filename: str, data: bytes) -> list[dict]:
    rows: list[list] = []
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        rows = [[("" if c is None else str(c)) for c in r] for r in ws.iter_rows(values_only=True)]
    else:
        text = data.decode("utf-8-sig", errors="replace")
        rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        return []
    headers = [HEADER_ALIASES.get(h.strip().lower(), "") for h in rows[0]]
    if "name_en" not in headers or "mobile" not in headers:
        raise HTTPException(400, "File needs at least 'name' and 'mobile' columns")
    out = []
    for r in rows[1:]:
        if not any(str(c).strip() for c in r):
            continue
        out.append({h: str(v).strip() for h, v in zip(headers, r) if h})
    return out


def bulk_import(session: Session, event: Event, rows: list[dict]) -> tuple[int, list[str]]:
    created, errors = 0, []
    for i, row in enumerate(rows, start=2):
        errs = validate(row)
        if errs:
            errors.append(f"Row {i}: {'; '.join(errs)}")
            continue
        try:
            create_registration(session, event, row, approved=True)
            created += 1
        except HTTPException as e:
            errors.append(f"Row {i}: {e.detail}")
    return created, errors
